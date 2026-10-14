import os
import re
from typing import Dict, Any, List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.config import settings

def sanitize_filename(name: str) -> str:
    return re.sub(r'[^\w\-_\.]', '_', name)

class ResumeGenerator:
    @staticmethod
    def generate_pdf(resume_data: Dict[str, Any], user_profile: Dict[str, Any], output_filename: str) -> str:
        """
        Builds a crisp, single/multi-page ATS-optimized PDF resume using ReportLab.
        """
        output_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            leftMargin=40,
            rightMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()
        
        # Custom Typography
        name_style = ParagraphStyle(
            'NameStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#1E293B')
        )

        contact_style = ParagraphStyle(
            'ContactStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#475569')
        )

        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=15,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=10,
            spaceAfter=4,
            textTransform='uppercase'
        )

        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            alignment=TA_JUSTIFY,
            textColor=colors.HexColor('#334155')
        )

        job_title_style = ParagraphStyle(
            'JobTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1E293B')
        )

        job_sub_style = ParagraphStyle(
            'JobSub',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#64748B')
        )

        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            leftIndent=14,
            firstLineIndent=-10,
            textColor=colors.HexColor('#334155')
        )

        story = []

        # 1. Header (Name & Contact)
        candidate_name = user_profile.get("full_name") or user_profile.get("user", {}).get("full_name") or "Candidate Name"
        story.append(Paragraph(candidate_name, name_style))
        story.append(Spacer(1, 4))

        contact_bits = []
        if user_profile.get("email"):
            contact_bits.append(user_profile["email"])
        if user_profile.get("phone"):
            contact_bits.append(user_profile["phone"])
        if user_profile.get("current_city"):
            contact_bits.append(user_profile["current_city"])
        if user_profile.get("linkedin_url"):
            contact_bits.append("LinkedIn")
        if user_profile.get("github_url"):
            contact_bits.append("GitHub")

        contact_text = " • ".join(contact_bits) if contact_bits else "candidate@example.com • Portfolio"
        story.append(Paragraph(contact_text, contact_style))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=8))

        # 2. Professional Summary
        summary = resume_data.get("summary")
        if summary:
            story.append(Paragraph("Professional Summary", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=4))
            story.append(Paragraph(summary, body_style))
            story.append(Spacer(1, 8))

        # 3. Core Skills
        skills_sec = resume_data.get("skills_section", {})
        if skills_sec:
            story.append(Paragraph("Skills & Competencies", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=4))
            
            if isinstance(skills_sec, dict):
                for cat, sk_list in skills_sec.items():
                    if sk_list and isinstance(sk_list, list):
                        line = f"<b>{cat}:</b> " + ", ".join([str(s) for s in sk_list])
                        story.append(Paragraph(line, body_style))
                        story.append(Spacer(1, 2))
            elif isinstance(skills_sec, list):
                line = "<b>Technical Skills:</b> " + ", ".join([str(s) for s in skills_sec])
                story.append(Paragraph(line, body_style))
            story.append(Spacer(1, 6))

        # 4. Professional Experience
        experiences = resume_data.get("experience_section", [])
        if experiences:
            story.append(Paragraph("Professional Experience", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=4))
            
            for exp in experiences:
                comp = exp.get("company", "Company")
                pos = exp.get("position", "Position")
                dates = f"{exp.get('start_date', '')} – {exp.get('end_date', 'Present')}"
                loc = exp.get("location", "")
                
                header_table = Table(
                    [[Paragraph(f"<b>{pos}</b> – {comp}", job_title_style), Paragraph(dates, job_sub_style)]],
                    colWidths=[380, 150]
                )
                header_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('ALIGN', (1,0), (1,0), 'RIGHT'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ]))
                story.append(header_table)

                if exp.get("description"):
                    story.append(Paragraph(exp["description"], body_style))
                    story.append(Spacer(1, 2))

                for bullet in exp.get("bullet_points", []):
                    story.append(Paragraph(f"• {bullet}", bullet_style))
                    story.append(Spacer(1, 1.5))
                story.append(Spacer(1, 6))

        # 5. Education
        educations = resume_data.get("education_section", [])
        if educations:
            story.append(Paragraph("Education", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=4))
            for edu in educations:
                inst = edu.get("institution", "University")
                deg = edu.get("degree", "Degree")
                fos = edu.get("field_of_study", "")
                dates = f"{edu.get('start_date', '')} – {edu.get('end_date', '')}".strip(" – ")
                deg_line = f"<b>{deg}</b>" + (f" in {fos}" if fos else "")
                
                edu_table = Table(
                    [[Paragraph(deg_line, job_title_style), Paragraph(dates, job_sub_style)]],
                    colWidths=[380, 150]
                )
                edu_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('ALIGN', (1,0), (1,0), 'RIGHT'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                ]))
                story.append(edu_table)
                story.append(Paragraph(inst, body_style))
                story.append(Spacer(1, 4))

        # 6. Projects (if present)
        projects = resume_data.get("projects_section", [])
        if projects:
            story.append(Paragraph("Key Projects", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=4))
            for proj in projects[:3]:
                title = proj.get("title", "Project")
                techs = ", ".join(proj.get("technologies", []))
                story.append(Paragraph(f"<b>{title}</b>" + (f" | <i>{techs}</i>" if techs else ""), job_title_style))
                if proj.get("description"):
                    story.append(Paragraph(proj["description"], body_style))
                for ach in proj.get("key_achievements", []):
                    story.append(Paragraph(f"• {ach}", bullet_style))
                story.append(Spacer(1, 4))

        doc.build(story)
        return output_path

    @staticmethod
    def generate_docx(resume_data: Dict[str, Any], user_profile: Dict[str, Any], output_filename: str) -> str:
        """
        Builds a crisp, standard Word DOCX resume using python-docx.
        """
        output_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        doc = docx.Document()

        # Page Margins
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        # Name
        candidate_name = user_profile.get("full_name") or user_profile.get("user", {}).get("full_name") or "Candidate Name"
        name_para = doc.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name_para.add_run(candidate_name)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(30, 41, 59)

        # Contact Info
        contact_bits = []
        if user_profile.get("email"): contact_bits.append(user_profile["email"])
        if user_profile.get("phone"): contact_bits.append(user_profile["phone"])
        if user_profile.get("current_city"): contact_bits.append(user_profile["current_city"])
        if user_profile.get("linkedin_url"): contact_bits.append(user_profile["linkedin_url"])

        contact_para = doc.add_paragraph(" | ".join(contact_bits) if contact_bits else "candidate@example.com")
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in contact_para.runs:
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(100, 116, 139)

        # Summary
        summary = resume_data.get("summary")
        if summary:
            h = doc.add_heading(level=2)
            hrun = h.add_run("PROFESSIONAL SUMMARY")
            hrun.font.size = Pt(11)
            hrun.bold = True
            hrun.font.color.rgb = RGBColor(15, 23, 42)
            p = doc.add_paragraph(summary)
            p.paragraph_format.space_after = Pt(8)

        # Skills
        skills_sec = resume_data.get("skills_section", {})
        if skills_sec:
            h = doc.add_heading(level=2)
            hrun = h.add_run("SKILLS & COMPETENCIES")
            hrun.font.size = Pt(11)
            hrun.bold = True
            hrun.font.color.rgb = RGBColor(15, 23, 42)
            
            if isinstance(skills_sec, dict):
                for cat, sk_list in skills_sec.items():
                    if sk_list and isinstance(sk_list, list):
                        p = doc.add_paragraph()
                        r1 = p.add_run(f"{cat}: ")
                        r1.bold = True
                        p.add_run(", ".join([str(s) for s in sk_list]))
                        p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_after = Pt(8)

        # Experience
        experiences = resume_data.get("experience_section", [])
        if experiences:
            h = doc.add_heading(level=2)
            hrun = h.add_run("PROFESSIONAL EXPERIENCE")
            hrun.font.size = Pt(11)
            hrun.bold = True
            hrun.font.color.rgb = RGBColor(15, 23, 42)

            for exp in experiences:
                comp = exp.get("company", "Company")
                pos = exp.get("position", "Position")
                dates = f"{exp.get('start_date', '')} – {exp.get('end_date', 'Present')}"
                
                exp_header = doc.add_paragraph()
                r_pos = exp_header.add_run(f"{pos} | {comp}")
                r_pos.bold = True
                r_date = exp_header.add_run(f"    ({dates})")
                r_date.italic = True

                if exp.get("description"):
                    doc.add_paragraph(exp["description"])

                for bullet in exp.get("bullet_points", []):
                    bp = doc.add_paragraph(bullet, style='List Bullet')
                    bp.paragraph_format.space_after = Pt(2)
                exp_header.paragraph_format.space_after = Pt(4)

        # Education
        educations = resume_data.get("education_section", [])
        if educations:
            h = doc.add_heading(level=2)
            hrun = h.add_run("EDUCATION")
            hrun.font.size = Pt(11)
            hrun.bold = True
            hrun.font.color.rgb = RGBColor(15, 23, 42)

            for edu in educations:
                inst = edu.get("institution", "University")
                deg = edu.get("degree", "Degree")
                fos = edu.get("field_of_study", "")
                dates = f"{edu.get('start_date', '')} – {edu.get('end_date', '')}".strip(" – ")
                
                ep = doc.add_paragraph()
                r_deg = ep.add_run(f"{deg} in {fos}" if fos else deg)
                r_deg.bold = True
                ep.add_run(f" - {inst} ({dates})")
                ep.paragraph_format.space_after = Pt(4)

        doc.save(output_path)
        return output_path

resume_generator = ResumeGenerator()
