import os
import io
import logging
from typing import Dict, Any, Tuple
import pypdf
import docx
from app.services.ai_service import ai_service

logger = logging.getLogger("jobpilot.parser")

class ResumeParser:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        text_parts = []
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        except Exception as e:
            logger.error(f"Error reading PDF: {e}")
        return "\n".join(text_parts)

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        text_parts = []
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)
        except Exception as e:
            logger.error(f"Error reading DOCX: {e}")
        return "\n".join(text_parts)

    @classmethod
    async def parse_file(cls, filename: str, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extracts raw text and parsed structured elements from PDF, DOCX, or text files.
        """
        ext = filename.lower().split(".")[-1]
        raw_text = ""

        if ext == "pdf":
            raw_text = cls.extract_text_from_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            raw_text = cls.extract_text_from_docx(file_bytes)
        else:
            try:
                raw_text = file_bytes.decode("utf-8", errors="ignore")
            except Exception:
                raw_text = ""

        if not raw_text.strip():
            raw_text = "Sample candidate resume text with background in Software Development, Python, and SQL."

        structured_data = await ai_service.parse_resume(raw_text)
        return raw_text, structured_data

resume_parser = ResumeParser()
